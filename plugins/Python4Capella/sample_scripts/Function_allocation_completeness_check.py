from distutils.tests.test_register import Inputs

# include needed for the Capella modeller API
include('workspace://Python4Capella/simplified_api/capella.py')
if False:
    from simplified_api.capella import *
    
# include needed for utilities
include('workspace://Python4Capella/utilities/CapellaPlatform.py')
if False:
    from utilities.CapellaPlatform import *
    
# include needed to read/write xlsx files
from openpyxl import *
from openpyxl.worksheet.table import Table, TableStyleInfo

# change this path to execute the script on your model (here is the IFE sample). 
# comment it if you want to use the "Run configuration" instead
aird_path = '/In-Flight Entertainment System/In-Flight Entertainment System.aird'

'''
#Here is the "Run Configuration" part to uncomment if you want to use this functionality :

#check parameter numbers
if len(argv) != 1:
    # use IFE default values
    aird_path = "/In-Flight Entertainment System/In-Flight Entertainment System.aird"
else:
    # Load the Capella model from the first argument of the script
    aird_path = argv[0]
'''


model = CapellaModel()
model.open(aird_path)

# gets the SystemEngineering and print its name
se = model.get_system_engineering()
print('starting export of model ' + se.get_name())

# preparing excel file export
project_name = aird_path[0:(aird_path.index("/", 1) + 1)]
project = CapellaPlatform.getProject(project_name)
folder = CapellaPlatform.getFolder(project, 'Exports - Python4Capella developments')
xlsx_file_name = CapellaPlatform.getAbsolutePath(folder) + '/' + 'KPI Physical Functions allocation test.xlsx'
# create a workbook
workbook = Workbook()

# writing excel file header
worksheet = workbook.active
worksheet.title = 'PF allocation'
worksheet["A1"] = 'Physical Function '
worksheet["B1"] = 'Allocating Component'
worksheet["C1"] = 'Allocating Actor'
worksheet["D1"] = 'Function Allocated?'


i=1
clf = 0
# retrieving elements from the model
all_F = se.get_all_contents_by_type(PhysicalFunction)
#all_FE = se.get_all_contents_by_type(FunctionalExchange)
#all_LC = se.get_all_contents_by_type(PhysicalComponent)
#all_LA = se.get_all_contents_by_type(PhysicalActor)

for f in all_F :
    leaf = 1
    cont_f = f.get_contained_physical_functions()
    for f2 in cont_f:
        leaf = 0
    if leaf == 1:
        i=i+1
        ac = f.get_allocating_component()
        worksheet["A" + str(i)] = f.get_name()
        if (isinstance(f.get_allocating_component(), PhysicalComponent)):
            worksheet["B" + str(i)] = ac.get_name()
            worksheet["D" + str(i)] = "OK"
        elif (isinstance(f.get_allocating_component(), PhysicalActor)):
            worksheet["C" + str(i)] = ac.get_name()
            worksheet["D" + str(i)] = "OK"
        else :
            worksheet["D" + str(i)] = "Function not allocated"


#Generating the table        
tab = Table(displayName="Table1", ref="A1:D"+str(i))

style = TableStyleInfo(name="TableStyleMedium1", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=False)
tab.tableStyleInfo = style

worksheet.add_table(tab)

#COUNTIF function for the KPIs calculation
worksheet["F2"].value = 'KPI'
worksheet["F3"].value = 'Number of Leaf Functions'
worksheet["F4"].value = 'Number Allocated'
worksheet["F5"].value = '% of Completeness'

worksheet["G2"].value = 'Progress'
worksheet["G3"].value = i-1
worksheet["G4"].value = '=COUNTIF(D:D, "OK")'
worksheet["G5"].value = '=G4/G3'
worksheet["G5"].number_format = '0%'

#Generating the table        
tab = Table(displayName="Table2", ref="F2:G5")

style = TableStyleInfo(name="TableStyleMedium1", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=False)
tab.tableStyleInfo = style

worksheet.add_table(tab)

worksheet.column_dimensions["A"].width = 40
worksheet.column_dimensions["B"].width = 20
worksheet.column_dimensions["C"].width = 20
worksheet.column_dimensions["D"].width = 20
worksheet.column_dimensions["F"].width = 25
worksheet.column_dimensions["G"].width = 15
                   
# Save the xlsx file
workbook.save(xlsx_file_name)

print('saving excel file')

# refresh 
CapellaPlatform.refresh(folder)